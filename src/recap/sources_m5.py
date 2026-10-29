"""Bounded document dependencies, real page renders and native DOCX cell arithmetic.

No model calls, target access, OCR, inferred PDF cells, or clinical certification.
"""
from __future__ import annotations

from copy import deepcopy
import hashlib
from pathlib import Path
import re

from docx import Document
from PIL import Image
from pypdf import PdfReader

from .calculations import PercentageInput, check_reported_percentage
from .document_sources import HEADING, REF, index_document, render_pdf_page
from .io import sha256

CELL = re.compile(r'^table/(\d+)/row/(\d+)/cell/(\d+)$')
CAPTION = re.compile(r'^\s*((?:Table|Fig(?:ure)?\.?)\s+(?:[A-Z]?\d+(?:[.\-]\d+)?|[IVX]+))\b', re.I)
FOOTNOTE = re.compile(r'^\s*(?:Notes?\s*:|Footnotes?\s*:|[*†‡]\s*|[a-z]\)\s+)', re.I)
HEADERS = ('Population', 'Timepoint', 'Outcome', 'Unit', 'Calculation',
           'Event count (n)', 'Total participants (N)', 'Reported percent (%; one decimal place)')
SIMPLE_METHOD = 'Simple event proportion (n/N); unweighted; unadjusted'


def _text_hash(value):
    return hashlib.sha256(value.encode('utf-8')).hexdigest()


def _ref(value):
    return re.sub(r'\s+', ' ', re.sub(r'^fig(?:ure)?\.?', 'figure', value.strip(), flags=re.I)).lower()


def _candidate_dict(candidate):
    return candidate.model_dump() if hasattr(candidate, 'model_dump') else candidate


def build_dependency_index(index: dict) -> dict:
    """Index every available native source without changing its text, ID or location.

    Line/section/caption/footnote recognition is a declared heuristic, not complete
    semantic recovery. A PDF line offset is never represented as a bounding box.
    """
    result = deepcopy(index)
    blocks, sections, captions, footnotes, references, cells = [], [], [], [], [], []
    current_section = None
    for source in result['sources']:
        sid, value = source['source_id'], source['text']
        source['text_sha256'] = _text_hash(value)
        offset = 0
        source_sections = []
        for line in value.splitlines(keepends=True):
            raw = line.rstrip('\r\n')
            entry = {'source_id': sid, 'raw_start': offset, 'raw_end': offset + len(raw),
                     'text_sha256': _text_hash(raw), 'location': source['location']}
            if raw.strip():
                blocks.append(entry)
            if HEADING.fullmatch(raw.strip()):
                current_section = raw.strip()
                source_sections.append(current_section)
                sections.append(entry | {'heading': current_section, 'detection': 'declared native line heading'})
            if match := CAPTION.match(raw):
                captions.append(entry | {'reference': _ref(match[1]), 'detection': 'line-start reference; caption status unverified'})
            if FOOTNOTE.match(raw):
                footnotes.append(entry | {'detection': 'explicit note/footnote line marker'})
            offset += len(line)
        source['m5_section_hint'] = current_section
        source['m5_sections_on_source'] = source_sections
        for match in REF.finditer(value):
            references.append({'source_id': sid, 'reference': _ref(match[0]), 'raw_start': match.start(), 'raw_end': match.end()})
        if match := CELL.fullmatch(source['location']):
            cell = {'source_id': sid, 'table': int(match[1]), 'row': int(match[2]), 'column': int(match[3])}
            source['m5_cell'] = cell
            cells.append(cell)
    result['inventory']['m5_dependencies'] = {
        'rule_version': 'M5_native_dependencies_v1', 'blocks': blocks, 'sections': sections,
        'caption_candidates': captions, 'footnote_candidates': footnotes, 'references': references,
        'native_cells': cells, 'all_available_sources_indexed': True,
        'semantic_inventory_complete': False, 'pdf_cell_or_bbox_mapping': 'unverified',
    }
    return result


def _dependency_order(index, candidates):
    index = index if 'm5_dependencies' in index['inventory'] else build_dependency_index(index)
    lookup = {s['source_id']: s for s in index['sources']}
    if len(lookup) != len(index['sources']):
        raise ValueError('duplicate source IDs')
    candidates = [_candidate_dict(c) for c in candidates]
    required = list(dict.fromkeys(q['source_id'] for c in candidates for q in c['sources']))
    if any(sid not in lookup for sid in required):
        raise ValueError('candidate source outside document')
    deps = index['inventory']['m5_dependencies']
    reasons = {sid: ['candidate_evidence'] for sid in required}
    mandatory = list(required)
    for sid in required:
        if cell := lookup[sid].get('m5_cell'):
            for entry in deps['native_cells']:
                if entry['table'] == cell['table'] and entry['row'] in (0, cell['row']):
                    if entry['source_id'] not in mandatory:
                        mandatory.append(entry['source_id'])
                    reasons.setdefault(entry['source_id'], []).append('same_native_table_row_or_header')
    text = ' '.join(str(c.get('description', '')) + ' ' + ' '.join(q['quote'] for q in c['sources']) for c in candidates)
    refs = {_ref(m[0]) for m in REF.finditer(text)}
    # Candidate pages may contain an explicit table/figure reference omitted from a short quote.
    refs |= {r['reference'] for r in deps['references'] if r['source_id'] in required}
    linked = []
    for caption in deps['caption_candidates']:
        if caption['reference'] in refs:
            sid = caption['source_id']
            linked.append(sid)
            reasons.setdefault(sid, []).append('referenced_caption_candidate')
    # Same-page notes are already included for PDF. Native DOCX note paragraphs
    # adjacent to a referenced caption/table remain explicit dependency candidates.
    positions = {s['source_id']: i for i, s in enumerate(index['sources'])}
    referenced_tables = {lookup[sid]['m5_cell']['table'] for sid in mandatory if lookup[sid].get('m5_cell')}
    table_positions = [positions[s['source_id']] for s in index['sources'] if s.get('m5_cell', {}).get('table') in referenced_tables]
    for footnote in deps['footnote_candidates']:
        sid = footnote['source_id']
        if sid in linked or any(abs(positions[sid]-p) <= 1 for p in table_positions):
            linked.append(sid)
            reasons.setdefault(sid, []).append('explicit_adjacent_footnote')
    methods_results = []
    for s in index['sources']:
        hint = ' '.join(s.get('m5_sections_on_source', []) + [s.get('m5_section_hint') or ''])
        if re.search(r'\b(methods|results)\b', hint, re.I):
            methods_results.append(s['source_id'])
            reasons.setdefault(s['source_id'], []).append('methods_or_results_section')
    return index, list(dict.fromkeys(mandatory + linked + methods_results)), set(mandatory), reasons, set(linked)


def select_dependency_sources(index: dict, candidates, *, max_chars=30000, max_sources=40) -> dict:
    """Preserve complete sources; fail rather than cut mandatory candidate/row cells."""
    if max_chars < 1 or max_sources < 1:
        raise ValueError('positive source limits required')
    index, order, mandatory, reasons, linked = _dependency_order(index, candidates)
    lookup = {s['source_id']: s for s in index['sources']}
    if len(mandatory) > max_sources or sum(len(lookup[s]['text']) for s in mandatory) > max_chars:
        raise ValueError('complete required evidence exceeds source budget; no middle truncation')
    chosen, used = [], 0
    for sid in order:
        source = lookup[sid]
        if source['text'].strip() and len(chosen) < max_sources and used + len(source['text']) <= max_chars:
            chosen.append(source)
            used += len(source['text'])
    selected_ids = {s['source_id'] for s in chosen}
    if not mandatory <= selected_ids:
        raise ValueError('required evidence unavailable')
    return {'sources': chosen, 'coverage': {
        'selection_rule': 'M5 candidate evidence, complete native row/header, reference captions/notes, Methods/Results; whole sources only',
        'selected_sources': len(chosen), 'indexed_sources': len(lookup), 'selected_chars': used,
        'max_chars': max_chars, 'max_sources': max_sources, 'truncated_source_ids': [],
        'not_selected_source_ids': [sid for sid in lookup if sid not in selected_ids],
        'omitted_dependency_source_ids': sorted(linked - selected_ids),
        'required_sources_complete': True, 'dependency_candidates_complete': linked <= selected_ids,
        'selection_reasons': {sid: reasons[sid] for sid in selected_ids},
        'all_indexed_text_supplied': selected_ids == set(lookup), 'whole_document_review_complete': False,
        'new_global_discovery_performed': False,
    }}


def select_render_pages(index: dict, candidates, *, max_pages=2) -> dict:
    if index['condition'] != 'pdf_only' or max_pages not in (1, 2):
        raise ValueError('PDF-only actual page render limit is one or two')
    index, order, mandatory, reasons, linked = _dependency_order(index, candidates)
    lookup = {s['source_id']: s for s in index['sources']}
    # Caption pages first when linked, then directly quoted pages. Never default first-N.
    eligible = list(dict.fromkeys([s for s in order if s in linked] + [s for s in order if s in mandatory]))
    pages, records = [], []
    for sid in eligible:
        source = lookup[sid]
        if source['location_kind'] != 'page' or source['source_modality'] != 'pdf_native_text':
            raise ValueError('render candidate lacks native PDF page provenance')
        if source['page'] not in pages and len(pages) < max_pages:
            pages.append(source['page'])
            records.append({'page': source['page'], 'source_id': sid, 'reasons': reasons[sid]})
    return {'pages': pages, 'reasons': records, 'max_pages': max_pages,
            'not_rendered_pages': [p['page'] for p in index['inventory']['pages'] if p['page'] not in pages],
            'selection_uses_candidate_evidence_not_gold': True, 'whole_document_visual_review': False}


def render_referenced_pages(index: dict, candidates, output_dir: Path, *, dpi=120, max_pages=2) -> dict:
    path = Path(index['local_path'])
    if path.is_symlink() or sha256(path) != index['source_asset']['sha256']:
        raise ValueError('original PDF asset changed')
    selection = select_render_pages(index, candidates, max_pages=max_pages)
    reader = PdfReader(path)
    labels = reader.page_labels if '/PageLabels' in reader.trailer['/Root'] else None
    images = []
    for page in selection['pages']:
        output = Path(output_dir) / f'page_{page:04d}.png'
        if output.exists():
            raise FileExistsError('render output already exists')
        item = render_pdf_page(path, page, output, dpi=dpi)
        with Image.open(output) as image:
            width, height = image.size
        item.update(page_index_zero_based=page-1, page_ordinal_one_based=page,
                    pdf_page_label=labels[page-1] if labels else None,
                    printed_page_label='unverified', width=width, height=height,
                    source_asset=deepcopy(index['source_asset']),
                    native_page_source_id=next(s['source_id'] for s in index['sources'] if s.get('page') == page),
                    renderer_output_verified=True, codex_visual_inspection='NOT_RUN',
                    model_visual_observation='NOT_RUN', visual_content_accuracy='unverified')
        if item['parent_sha256'] != index['source_asset']['sha256']:
            raise ValueError('render parent hash mismatch')
        images.append(item)
    return {'images': images, 'selection': selection, 'ocr': 'NOT_RUN'}


def prepare_docx_derived(parent_path: Path, new_path: Path) -> dict:
    """Numbers occur only in separate cells; expected outcomes are test-only."""
    parent_path, new_path = Path(parent_path), Path(new_path)
    parent_hash = sha256(parent_path)
    if new_path.exists() or new_path.resolve() == parent_path.resolve():
        raise FileExistsError('derived fixture must use a fresh path')
    doc = Document()
    doc.add_heading('Cohort follow-up results', 0)
    doc.add_heading('Methods', 1)
    doc.add_paragraph('This descriptive cohort study reports simple event proportions. Each table row identifies its population, timepoint, outcome, participant unit and calculation method. Counts and percentages are reported only in table cells.')
    doc.add_heading('Results', 1)
    doc.add_paragraph('Table 1. Participant event outcomes by population and timepoint.')
    table = doc.add_table(rows=1, cols=len(HEADERS))
    for cell, label in zip(table.rows[0].cells, HEADERS):
        cell.text = label
    for values in [('Cohort A', 'Day 30', 'Event Alpha', 'participants', SIMPLE_METHOD, '30', '200', '15.0%'),
                   ('Cohort A', 'Day 30', 'Event Beta', 'participants', SIMPLE_METHOD, '30', '200', '20.0%'),
                   ('Cohort B', 'Day 30', 'Event Gamma', 'participants', SIMPLE_METHOD, '12', 'Not reported', '10.0%')]:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    doc.add_paragraph('Note: Within each row, the event count and total participants refer to the same population, timepoint and outcome. The reported percentage is the simple event count divided by total participants, expressed as a percent and rounded to one decimal place. A missing total remains unknown.')
    new_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(new_path)
    if sha256(parent_path) != parent_hash:
        raise ValueError('parent fixture changed')
    return {'self_authored_fixture': True, 'representative_medical_gold': False,
            'derived_from_path': str(parent_path), 'parent_sha256': parent_hash,
            'path': str(new_path), 'source_sha256': sha256(new_path), 'license': 'CC0-1.0',
            'changes': 'Separate native table cells with explicit row headers; original numeric prose removed in this new fixture only.',
            'expected_outcomes_embedded': False, 'model_calls': 0}


def execute_table_percentage(request, sources) -> dict:
    """Bind existing typed numeric evidence to complete native table rows/headers.

    This is a deliberately small English header/method grammar. Location checking,
    declared semantics and the model's assertions remain separately recorded.
    """
    lookup = {s['source_id']: s for s in sources}
    if len(lookup) != len(sources):
        raise ValueError('duplicate source IDs')
    numeric = [request.numerator, request.denominator, request.reported_percent]
    args = {'numerator': request.numerator.value, 'denominator': request.denominator.value if request.denominator else '',
            'reported_percent': request.reported_percent.value.removesuffix('%').strip(), 'decimal_places': request.decimal_places}
    located = []
    literal_checks = []
    for pos, quote in enumerate(numeric):
        source = lookup.get(quote.source_id) if quote else None
        match = CELL.fullmatch(source['location']) if source else None
        located.append(tuple(map(int, match.groups())) if match and source['source_modality'] == 'docx_native_text' else None)
        pattern = r'\d+(?:\.\d+)?\s*%?' if pos == 2 else r'\d+'
        literal = source['text'].strip() if source else ''
        expected = quote.value.removesuffix('%').strip() if quote else ''
        literal_checks.append(bool(quote and source and quote.quote.strip() == literal and re.fullmatch(pattern, literal)
                                   and literal.removesuffix('%').strip() == expected
                                   and source.get('grid_span') == 1 and source.get('merged_origin') == source['location']))
    same_row = all(located) and len({x[:2] for x in located}) == 1
    valid_row = bool(same_row and located[0][1] > 0)
    rows = {}
    if valid_row:
        table, row = located[0][:2]
        for source in sources:
            if (match := CELL.fullmatch(source['location'])) and int(match[1]) == table and int(match[2]) in (0, row):
                key = (int(match[2]), int(match[3]))
                if key in rows:
                    raise ValueError('duplicate native table location')
                rows[key] = source
        def value(r, c):
            return rows.get((r, c), {}).get('text', '').strip()
        headers_valid = all(value(0, col) == text for col, text in enumerate(HEADERS))
        roles_valid = [x[2] for x in located] == [5, 6, 7]
        population = bool(re.fullmatch(r'(?:Cohort|Group|Arm) [A-Za-z0-9_-]+', value(row, 0))) and not re.search(r'unknown|unverified|unspecified', value(row, 0), re.I)
        timepoint = bool(re.fullmatch(r'(?:Day|Week|Month|Year) \d+|Baseline', value(row, 1)))
        outcome = bool(re.fullmatch(r'Event [A-Za-z][A-Za-z0-9 _-]*', value(row, 2))) and not re.search(r'unknown|not reported|unverified', value(row, 2), re.I)
        unit = value(row, 3) == 'participants'
        simple = value(row, 4) == SIMPLE_METHOD
        binding_sources = [s for _, s in sorted(rows.items())]
    else:
        headers_valid = roles_valid = population = timepoint = outcome = unit = simple = False
        binding_sources = []
    source = lookup.get(request.context.source_id)
    context_valid = bool(source and request.context.quote in source['text'])
    # A context from a different native row/table cannot substitute for the bound row.
    context_location = CELL.fullmatch(source['location']) if source else None
    if context_location and valid_row:
        context_valid &= int(context_location[1]) == table and int(context_location[2]) in (0, row)
    decimals = len(args['reported_percent'].split('.')[1]) if '.' in args['reported_percent'] else 0
    checks = {'exact_cell_literals': all(literal_checks), 'same_native_table_row': bool(valid_row),
              'binding_native_cells_unmerged': bool(binding_sources) and all(s.get('grid_span') == 1 and s.get('merged_origin') == s['location'] and s.get('source_modality') == 'docx_native_text' for s in binding_sources),
              'single_original_asset': bool(binding_sources) and len({s['asset_id'] for s in binding_sources}) == 1,
              'explicit_headers': bool(headers_valid), 'numeric_column_roles': bool(roles_valid),
              'population_explicit': bool(population), 'timepoint_explicit': bool(timepoint),
              'outcome_explicit': bool(outcome), 'participant_unit_explicit': bool(unit),
              'simple_proportion_explicit': bool(simple), 'context_quote_in_scope': context_valid,
              'rounding_declared_and_matches': decimals == request.decimal_places == 1}
    verified = all(checks.values())
    args.update(same_population_verified=verified and request.same_population_claimed,
                same_timepoint_verified=verified and request.same_timepoint_claimed,
                simple_proportion_verified=verified and request.simple_proportion_claimed)
    result = check_reported_percentage(PercentageInput.model_validate(args))
    return {'tool': 'check_reported_percentage', 'tool_version': sha256(Path(__file__).with_name('calculations.py')),
            'arguments': args, 'result': result, 'executed': True,
            'numeric_sources': [q.model_dump() if q else None for q in numeric],
            'host_checks': checks, 'host_binding': {'method': 'M5_native_DOCX_row_header_v1',
                'sources': [{'source_id': s['source_id'], 'location': s['location'], 'text': s['text'],
                             'text_sha256': _text_hash(s['text'])} for s in binding_sources]},
            'model_assertions': {'same_population': request.same_population_claimed,
                'same_timepoint': request.same_timepoint_claimed, 'simple_proportion': request.simple_proportion_claimed},
            'execution_scope': 'native_structure_and_explicit_declared_conditions_arithmetic_only',
            'statistical_and_clinical_validity': 'unverified',
            'applicability': 'English native DOCX explicit participant event proportions; no PDF cell or arbitrary-design certification'}
