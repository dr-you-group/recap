"""Observed locations only: PDF pages, JATS paths, and DOCX cells."""
from __future__ import annotations

from pathlib import Path
from lxml import etree
from pypdf import PdfReader
from docx import Document

from .download import validate_file, xml_root
from .io import sha256


def text(node) -> str:
    return " ".join(" ".join(node.itertext()).split()) if node is not None else ""


def extract_pdf(path: Path) -> dict:
    validate_file(path, "pdf")
    reader = PdfReader(path)
    pages = [{"page": i + 1, "text": page.extract_text() or "",
              "method": "pypdf_native", "bbox": None} for i, page in enumerate(reader.pages)]
    return {"format": "pdf", "source_sha256": sha256(path), "pages": pages,
            "empty_pages": [x["page"] for x in pages if not x["text"].strip()],
            "ocr": "NOT_RUN", "cell_mapping": "unverified"}


def extract_jats(path: Path) -> dict:
    root = xml_root(path.read_bytes())
    if etree.QName(root).localname != "article":
        raise ValueError("expected a JATS article")
    tree = root.getroottree()
    identifiers = {x.get("pub-id-type"): text(x) for x in root.xpath("./front/article-meta/article-id")}
    tables = []
    for table in root.xpath(".//table-wrap"):
        rows = []
        for row in table.xpath(".//tr"):
            rows.append([{"text": text(c), "rowspan": c.get("rowspan", "1"),
                          "colspan": c.get("colspan", "1"), "xpath": tree.getpath(c)}
                         for c in row if etree.QName(c).localname in {"th", "td"}])
        tables.append({"id": table.get("id"), "label": text(table.find("label")),
                       "caption": text(table.find("caption")), "rows": rows,
                       "footnotes": [text(x) for x in table.xpath("./table-wrap-foot//fn")],
                       "xpath": tree.getpath(table), "pdf_mapping": "unverified"})
    figures = []
    for figure in root.xpath(".//fig"):
        figures.append({"id": figure.get("id"), "label": text(figure.find("label")),
                        "caption": text(figure.find("caption")), "xpath": tree.getpath(figure),
                        "hrefs": [g.get("{http://www.w3.org/1999/xlink}href") for g in figure.xpath(".//graphic")],
                        "rights_notices": [text(x) for x in figure.xpath(".//permissions | .//copyright-statement | .//license")]})
    licenses = [{"text": text(x), "href": x.get("{http://www.w3.org/1999/xlink}href"),
                 "xpath": tree.getpath(x)} for x in root.xpath("./front/article-meta/permissions/license")]
    return {"format": "jats", "source_sha256": sha256(path), "identifiers": identifiers,
            "title": text(root.find("./front/article-meta/title-group/article-title")),
            "language": root.get("{http://www.w3.org/XML/1998/namespace}lang"),
            "article_type": root.get("article-type"), "licenses": licenses,
            "tables": tables, "figures": figures,
            "figure_references": [{"rid": x.get("rid"), "text": text(x), "xpath": tree.getpath(x)}
                                  for x in root.xpath(".//xref[@ref-type='fig']")],
            "table_references": [{"rid": x.get("rid"), "text": text(x), "xpath": tree.getpath(x)}
                                 for x in root.xpath(".//xref[@ref-type='table']")],
            "review_gold": False}


def extract_docx(path: Path) -> dict:
    doc = Document(path)
    tables = []
    for ti, table in enumerate(doc.tables):
        seen = {}
        rows = []
        for ri, row in enumerate(table.rows):
            cells = []
            for ci, cell in enumerate(row.cells):
                location = f"table/{ti}/row/{ri}/cell/{ci}"
                canonical = seen.setdefault(cell._tc, location)
                cells.append({"text": cell.text, "location": location,
                              "merged_origin": canonical, "grid_span": cell._tc.grid_span})
            rows.append(cells)
        tables.append({"rows": rows})
    return {"format": "docx", "source_sha256": sha256(path),
            "paragraphs": [p.text for p in doc.paragraphs], "tables": tables,
            "pdf_mapping": "unverified", "fixture_is_submission_representative": False}


def extract_file(path: Path) -> dict:
    return {".pdf": extract_pdf, ".xml": extract_jats, ".docx": extract_docx}[path.suffix.lower()](path)
