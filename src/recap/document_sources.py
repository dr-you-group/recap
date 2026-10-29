"""Host-owned document source inventory; no targets, model calls, or inferred coordinates."""
from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path
import random
import re
import shutil
import subprocess
from collections import Counter

from docx import Document
from lxml import etree
from pypdf import PdfReader

from .contracts import Asset
from .download import xml_root
from .extract import extract_docx, extract_jats, extract_pdf, text
from .io import canonical, now, read_jsonl, sha256, write_json

CONDITIONS = {'pdf_only': {'.pdf'}, 'docx_native': {'.docx'}, 'xml_assisted': {'.xml', '.pdf'}}
MEDIA = {'.pdf': 'application/pdf', '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', '.xml': 'application/xml'}
MODALITY = {'.pdf': 'pdf_native_text', '.docx': 'docx_native_text', '.xml': 'jats_xml_text'}
REF = re.compile(r'\b(?:Table|Fig(?:ure)?\.?)\s+(?:[A-Z]?\d+(?:[.\-]\d+)?|[IVX]+)\b', re.I)
HEADING = re.compile(r'^(?:\d+(?:\.\d+)*[. ]+)?(abstract|introduction|background|materials and methods|methods|results|discussion|conclusions?|references|limitations|acknowledgments|funding|supplementary materials?)\s*$', re.I)


def _id(prefix, *values):
    return prefix + '_' + hashlib.sha256(canonical(list(values))).hexdigest()[:32]


def _asset(path, provenance):
    checksum = sha256(path)
    if provenance:
        allowed = {k:v for k,v in provenance.items() if k in Asset.model_fields}
        original = Asset.model_validate(allowed)
        if original.sha256 != checksum or original.media_type != MEDIA[path.suffix.lower()]:
            raise ValueError('document provenance hash or media type mismatch')
        base = original.model_dump()
    else:
        base = {'source_url': 'local_document', 'object_key': None, 'source_identity': {},
                'version': checksum, 'path': str(path), 'media_type': MEDIA[path.suffix.lower()],
                'sha256': checksum, 'size': path.stat().st_size, 'retrieved_at': now(),
                'license': 'user_provided_unverified', 'license_evidence': 'No license inferred from local possession',
                'rights_status': 'unresolved', 'work_family': None, 'changes': 'Original local bytes; no alteration'}
    base['asset_id'] = _id('a',checksum)
    return Asset.model_validate(base).model_dump()


def _source(asset, modality, kind, location, value, **extra):
    return {'source_id': _id('e',asset['sha256'],kind,location), 'asset_id':asset['asset_id'],
            'source_modality':modality, 'location_kind':kind, 'location':location,
            'text':value, **extra}


def _refs(sources):
    return [{'source_id':s['source_id'],'location':s['location'],'reference':m.group(),
             'text_start':m.start(),'text_end':m.end()}
            for s in sources for m in REF.finditer(s['text'])]


def _pdf(path, asset):
    failed=[]
    try:
        extracted=extract_pdf(path)
        pages=extracted['pages']
    except Exception as error:
        # Only on extraction failure, preserve independent per-page failure evidence.
        reader=PdfReader(path); pages=[]
        for i,page in enumerate(reader.pages,1):
            try:value=page.extract_text() or ''
            except Exception as page_error:
                value='';failed.append({'location':f'page/{i}','reason':'native_parse_error','error_type':type(page_error).__name__})
            pages.append({'page':i,'text':value,'method':'pypdf_native','bbox':None})
        if not pages:failed.append({'location':'document','reason':'no_pages','error_type':type(error).__name__})
    sources=[_source(asset,'pdf_native_text','page',f"page/{p['page']}",p['text'],page=p['page'],
                     extraction_method='pypdf_native',bbox=None) for p in pages]
    for s in sources:
        if not s['text'].strip() and not any(x['location']==s['location'] for x in failed):
            failed.append({'source_id':s['source_id'],'location':s['location'],'reason':'empty_native_text','ocr':'NOT_RUN'})
    sections=[{'source_id':s['source_id'],'location':s['location'],'heading':line.strip(),
               'detection':'native_line_heading_heuristic'} for s in sources for line in s['text'].splitlines()
              if HEADING.fullmatch(line.strip())]
    inventory={'pages':[{'page':s['page'],'source_id':s['source_id'],'native_text_chars':len(s['text'])} for s in sources],
               'sections':sections,'tables':[],'figures':[],
               'table_cell_mapping':'unverified','figure_bbox_mapping':'unverified',
               'section_inventory_limit':'Only headings matching the declared native-line rule; complete semantic section recognition is unverified.',
               'visual_regions':'Not independently enumerated or read; native page text only. Rendered pages must be separately requested.'}
    return sources,inventory,failed


def _docx(path,asset):
    extracted=extract_docx(path);doc=Document(path); sources=[];sections=[];tables=[]
    para_index=table_index=0
    for child in doc.element.body:
        local=etree.QName(child).localname
        if local=='p':
            p=doc.paragraphs[para_index];location=f'paragraph/{para_index}'
            s=_source(asset,'docx_native_text','docx_paragraph',location,extracted['paragraphs'][para_index],
                      paragraph=para_index,style=p.style.name if p.style else None,extraction_method='python_docx_native')
            sources.append(s)
            if p.style and (p.style.name.startswith('Heading') or p.style.name=='Title'):
                sections.append({'source_id':s['source_id'],'location':location,'heading':p.text,'style':p.style.name,'detection':'DOCX paragraph style'})
            para_index+=1
        elif local=='tbl':
            cells=[]
            for row in extracted['tables'][table_index]['rows']:
                for cell in row:
                    s=_source(asset,'docx_native_text','docx_cell',cell['location'],cell['text'],
                              merged_origin=cell['merged_origin'],grid_span=cell['grid_span'],extraction_method='python_docx_native')
                    sources.append(s);cells.append({'source_id':s['source_id'],'location':s['location'],'merged_origin':s['merged_origin']})
            tables.append({'table':table_index,'cells':cells});table_index+=1
    # These are actual package/body elements, not invented mapping of their visual extent.
    embedded=sum(1 for r in doc.part.rels.values() if 'image' in r.reltype)
    textboxes=len(doc.element.xpath('.//w:txbxContent'))
    footnotes=[r.reltype.rsplit('/',1)[-1] for r in doc.part.rels.values() if r.reltype.endswith(('/footnotes','/endnotes'))]
    headers_footers=[r.reltype.rsplit('/',1)[-1] for r in doc.part.rels.values() if r.reltype.endswith(('/header','/footer'))]
    failed=[]
    for kind,count in [('embedded_image',embedded),('text_box',textboxes),('notes_part',len(footnotes)),('header_footer_part',len(headers_footers))]:
        if count:failed.append({'location':'DOCX package','reason':kind+'_not_indexed','count':count,'status':'not_processed'})
    inventory={'pages':[],'page_mapping':'unverified_no_layout_engine','sections':sections,'paragraphs':para_index,
               'tables':tables,'figures':[],'embedded_image_count':embedded,'section_inventory_limit':'Only actual Title/Heading styles; unstyled semantic headings are not claimed identified.'}
    return sources,inventory,failed


def _jats(path,asset):
    extracted=extract_jats(path);root=xml_root(path.read_bytes());tree=root.getroottree();sources=[]
    # Paragraphs/section titles/table cells/figure captions remain XML-assisted only.
    for node in root.xpath('.//article-title | .//abstract//p | .//body//p | .//back//p | .//sec/title | .//table-wrap//th | .//table-wrap//td | .//fig/caption | .//table-wrap/caption | .//table-wrap-foot//fn'):
        sources.append(_source(asset,'jats_xml_text','jats_xpath',tree.getpath(node),text(node),extraction_method='lxml_jats'))
    # XPath unions may include nested paragraph and caption records, explicitly preserving their actual locations.
    sections=[{'source_id':s['source_id'],'location':s['location'],'heading':s['text'],'detection':'JATS sec/title'} for s in sources if s['location'].endswith('/title')]
    inventory={'pages':[],'page_mapping':'unverified','sections':sections,
               'tables':extracted['tables'],'figures':extracted['figures'],
               'native_figure_references':extracted['figure_references'],'native_table_references':extracted['table_references']}
    failed=[{'location':f['xpath'],'reason':'external_figure_asset_not_loaded','status':'not_processed'} for f in extracted['figures'] if f['hrefs']]
    return sources,inventory,failed


def index_document(path:Path, condition:str, *, provenance:dict|None=None)->dict:
    path=Path(path)
    if condition not in CONDITIONS or path.suffix.lower() not in CONDITIONS[condition]:
        raise ValueError('document format prohibited by input condition')
    if path.is_symlink():raise ValueError('symlink document prohibited')
    asset=_asset(path,provenance)
    sources,inventory,failed={'.pdf':_pdf,'.docx':_docx,'.xml':_jats}[path.suffix.lower()](path,asset)
    if len({s['source_id'] for s in sources})!=len(sources):raise ValueError('duplicate actual source location')
    inventory['table_figure_references']=_refs(sources)
    return {'schema_version':'M4_document_sources_v1','local_path':str(path.resolve()),'document_id':_id('d',asset['sha256']),
            'condition':condition,'source_asset':asset,'sources':sources,'inventory':inventory,'parse_failed':failed,
            'coverage':{'indexed_sources':len(sources),'native_text_chars':sum(len(s['text']) for s in sources),
                        'parse_failed_or_not_processed_regions':len(failed),'ocr':'NOT_RUN',
                        'all_semantic_claims_discovered':False,'whole_document_review_complete':False},
            'provenance_authority':'host original asset and parser; model source_kind is not authoritative'}


def select_sources(index:dict, max_chars:int=30000, max_sources:int=40)->dict:
    if max_chars<=0 or max_sources<=0:raise ValueError('positive source budgets required')
    chosen=[];omitted=[];used=0;truncations=[]
    # Document order only: no claim, score, gold, or label-driven page selection.
    for source in index['sources']:
        if not source['text'].strip():
            omitted.append(source['source_id']);continue
        remaining=max_chars-used
        if len(chosen)>=max_sources or remaining<=0:
            omitted.append(source['source_id']);continue
        item=dict(source)
        if len(item['text'])>remaining:
            item['text']=item['text'][:remaining]
            item['excerpt']={'start_char':0,'end_char':remaining,'original_chars':len(source['text']),'truncated':True}
            truncations.append(source['source_id'])
        chosen.append(item);used+=len(item['text'])
    return {'sources':chosen,'coverage':{'selection_rule':'Original document order with fixed character/source limits; no claim/gold-based selection',
            'max_chars':max_chars,'max_sources':max_sources,'selected_sources':len(chosen),'selected_chars':used,
            'indexed_sources':len(index['sources']),'not_selected_source_ids':omitted,'truncated_source_ids':truncations,
            'all_indexed_text_supplied':not omitted and not truncations,'whole_document_review_complete':False}}


def host_source(index:dict, source_id:str, *, observation_kind:str|None=None)->dict:
    source=next((s for s in index['sources'] if s['source_id']==source_id),None)
    if source is None:raise ValueError('source ID outside this document')
    asset=index['source_asset']
    if source['asset_id']!=asset['asset_id']:raise ValueError('host asset binding mismatch')
    return {k:source[k] for k in ['source_id','asset_id','source_modality','location_kind','location']} | {
        'asset_media_type':asset['media_type'],'asset_sha256':asset['sha256'],
        'observation_kind':observation_kind,'observation_kind_authority':'model assertion; separate from host source modality'}


class RendererUnavailable(RuntimeError):
    pass


def available_pdf_renderer():
    if importlib.util.find_spec('pypdfium2'):return 'pypdfium2'
    if importlib.util.find_spec('fitz'):return 'pymupdf'
    for name in ['pdftoppm','pdftocairo','mutool']:
        if shutil.which(name):return name
    return None


def render_pdf_page(path:Path,page:int,output:Path,*,dpi:int=120)->dict:
    path=Path(path);output=Path(output)
    if path.suffix.lower()!='.pdf' or type(page) is not int or page<1 or not 36<=dpi<=300:
        raise ValueError('invalid actual PDF page render request')
    if page>len(PdfReader(path).pages):raise ValueError('page outside PDF')
    renderer=available_pdf_renderer()
    if renderer is None:raise RendererUnavailable('No installed PDF renderer; native text remains available, visual inspection NOT_RUN')
    if output.suffix.lower()!='.png':raise ValueError('page render output must be PNG')
    output.parent.mkdir(parents=True,exist_ok=True)
    if renderer=='pypdfium2':
        import pypdfium2
        doc=pypdfium2.PdfDocument(str(path));p=doc[page-1];bitmap=p.render(scale=dpi/72)
        bitmap.to_pil().save(output);bitmap.close();p.close();doc.close()
    elif renderer=='pymupdf':
        import fitz
        with fitz.open(path) as doc:doc[page-1].get_pixmap(dpi=dpi).save(output)
    elif renderer in {'pdftoppm','pdftocairo'}:
        subprocess.run([shutil.which(renderer),'-f',str(page),'-l',str(page),'-r',str(dpi),'-png','-singlefile',str(path),str(output.with_suffix(''))],check=True,capture_output=True,timeout=60)
    else:
        subprocess.run([shutil.which('mutool'),'draw','-r',str(dpi),'-o',str(output),str(path),str(page)],check=True,capture_output=True,timeout=60)
    from PIL import Image
    with Image.open(output) as image:
        image.verify()
        if image.format!='PNG':raise ValueError('renderer did not produce PNG')
    return {'asset_id':_id('a',sha256(output)),'source_id':_id('e',sha256(path),'rendered_page',page,dpi),
            'source_modality':'pdf_page_render','media_type':'image/png','path':str(output),
            'sha256':sha256(output),'parent_sha256':sha256(path),'page':page,'dpi':dpi,
            'renderer':renderer,'location':f'page/{page}','crop_bbox':None,'actual_render':True}


def prepare_docx(path:Path):
    """Self-authored manuscript fixture. Expected outcomes live only in tests."""
    doc=Document();doc.add_heading('Cohort follow-up results',0)
    doc.add_heading('Methods',1)
    doc.add_paragraph('This is a descriptive cohort study. Cohort A contains 200 participants at day 30. Event counts are mutually exclusive participant counts, and percentages are simple proportions out of the stated cohort denominator, rounded to one decimal place.')
    doc.add_heading('Results',1)
    doc.add_paragraph('At day 30, in cohort A, 30 of 200 participants experienced event Alpha; the reported percentage is 15.0%.')
    doc.add_paragraph('At day 30, in cohort A, 30 of 200 participants experienced event Beta; the reported percentage is 20.0%.')
    doc.add_paragraph('At day 30, in cohort B, 12 participants experienced event Gamma; the reported percentage is 10.0%. The total number of cohort B participants is not reported.')
    doc.add_paragraph('Table 1 provides the corresponding descriptive results. No causal effect is estimated.')
    table=doc.add_table(rows=1,cols=5)
    for c,v in zip(table.rows[0].cells,['Cohort','Timepoint','Event count','Total participants','Reported percent']):c.text=v
    for vals in [('A','Day 30','30','200','15.0'),('A','Day 30','30','200','20.0'),('B','Day 30','12','Not reported','10.0')]:
        for c,v in zip(table.add_row().cells,vals):c.text=v
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);doc.save(path)
    return {'self_authored_fixture':True,'representative_medical_gold':False,'source_sha256':sha256(path),
            'filename_contains_expected_labels':False,'expected_results_location':'tests/test_document_sources_m4.py only'}


def prepare_m4_documents(root:Path,*,seed:int=20261014)->dict:
    root=Path(root);rows=read_jsonl(root/'data/pmc/pilot_manifest.jsonl');eligible=[]
    for row in rows:
        if row['status']!='available' or not row.get('correspondence',{}).get('checks',{}).get('pdf_has_native_text'):continue
        for asset in row['assets']:
            if asset.get('role')=='pdf' and asset['rights_status']=='verified' and (root/'data/pmc'/asset['path']).is_file():
                eligible.append((row,asset))
    eligible.sort(key=lambda item:(item[0]['pmcid'],item[0]['article_version']))
    if not eligible:raise ValueError('no existing rights-verified native-text PMC PDF')
    random.Random(seed).shuffle(eligible);row,asset=eligible[0]
    pdf=root/'data/pmc'/asset['path'];pdf_index=index_document(pdf,'pdf_only',provenance=asset)
    work=root/'artifacts/private/m4/documents';work.mkdir(parents=True,exist_ok=True)
    write_json(work/'pdf_index.json',pdf_index,immutable=True)
    filename='d_'+hashlib.sha256(b'M4 native manuscript fixture version1').hexdigest()[:32]+'.docx'
    fixture_path=root/'data/m4'/filename
    if fixture_path.exists():raise ValueError('M4 fixture already exists; do not regenerate')
    fixture=prepare_docx(fixture_path);docx_index=index_document(fixture_path,'docx_native')
    docx_index['source_asset'].update(source_identity={'self_authored_fixture':'true'},license='CC0-1.0',
        license_evidence='RECAP authored the synthetic fixture for this task; not clinical gold',rights_status='verified')
    write_json(work/'docx_index.json',docx_index,immutable=True)
    value={'seed':seed,'selection_rule':'Existing pilot manifest availability, verified rights and previously observed native parsing; sorted PMC ID/version then seeded shuffle; no prediction or gold inputs.',
           'eligible_pdf_count':len(eligible),'pdf':{'path':str(pdf.relative_to(root)),'pmcid':row['pmcid'],'article_version':row['article_version'],
                'source_sha256':asset['sha256'],'license':asset['license'],'condition':'pdf_only',
                'native_pages':len(pdf_index['inventory']['pages']),'index_path':str((work/'pdf_index.json').relative_to(root)),
                'parse_failed':pdf_index['parse_failed'],'review_gold':False},
           'docx':{'path':str(fixture_path.relative_to(root)),'condition':'docx_native','index_path':str((work/'docx_index.json').relative_to(root)),**fixture},
           'pdf_renderer':available_pdf_renderer(),'new_downloads':0,'new_model_calls':0,
           'final_pool_accessed':False}
    write_json(root/'artifacts/private/m4/document_selection.json',value,immutable=True)
    return value
